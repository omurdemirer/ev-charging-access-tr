"""Hecht (2022) Mendeley Data 'Summary data.xlsx' (doi:10.17632/ddv53zsf9m.1) dosyasini inceler.
Amac: saatlik VARIS (Arrivals) ve DOLULUK (Occupation) profillerinin hangi sayfalarda, hangi
kirilimda (guc sinifi AC/DC, hafta ici/sonu, saat/15 dk) verildigini bulmak.
Dosya buyuk (~244 MB): openpyxl read_only ile yalnizca sayfa adlari ve ilk satirlar okunur.
"""
import sys, os, re
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
F = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "veri", "kaynaklar",
                 "Hecht2022_MendeleyData", "Summary data.xlsx")
wb = openpyxl.load_workbook(F, read_only=True, data_only=True)
print("Sayfa sayisi:", len(wb.sheetnames))
for i, name in enumerate(wb.sheetnames):
    print("  %3d  %s" % (i, name))

hedef = [s for s in wb.sheetnames if re.search(r"arriv|occup|durat|ends|profile|hour|week", s, re.I)]
print("\nHedef sayfalar:", hedef)
for s in hedef[:12]:
    ws = wb[s]
    print("\n==========", s, "| boyut:", ws.max_row, "x", ws.max_column)
    for k, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [c for c in row if c is not None]
        if vals:
            print("   ", [str(v)[:22] for v in vals[:12]])
        if k >= 7:
            break
